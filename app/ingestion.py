"""File ingestion: extract text -> split into chunks -> embed -> store in pgvector."""
import csv
import io
from pathlib import Path
from uuid import UUID

from docx import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langsmith import traceable
from pypdf import PdfReader

from app.config import get_settings
from app.embeddings import get_embeddings
from app.vectorstore import add_chunks

SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".md", ".docx", ".csv"}


def extract_text(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file type '{ext}'. Allowed: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )

    if ext == ".pdf":
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages)

    if ext == ".docx":
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)

    text = data.decode("utf-8-sig", errors="replace")
    if ext == ".csv":
        # Turn each row into "column: value; column: value" so rows stay meaningful.
        rows = csv.DictReader(io.StringIO(text))
        return "\n".join(
            "; ".join(f"{k}: {v}" for k, v in row.items() if k) for row in rows
        )
    return text  # .txt / .md


def split_text(text: str) -> list[str]:
    s = get_settings()
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=s.chunk_size, chunk_overlap=s.chunk_overlap
    )
    return [c for c in splitter.split_text(text) if c.strip()]


@traceable(name="ingest_file", run_type="chain")
def ingest_file(session_id: UUID, filename: str, data: bytes) -> int:
    text = extract_text(filename, data)
    chunks = split_text(text)
    if not chunks:
        raise ValueError("No readable text found in the uploaded file.")
    vectors = get_embeddings().embed_documents(chunks)
    return add_chunks(session_id, filename, chunks, vectors)